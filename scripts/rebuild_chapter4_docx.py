#!/usr/bin/env python3
"""Build Chapter 4 implementation/results DOCX (master's thesis academic prose)."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "docs" / "CHAPTER_4_IMPLEMENTATION_RESULTS.docx"
FIGURE = REPO / "experiments" / "figures" / "cer_comparison.png"
EXPERIMENTS = REPO / "experiments"


def load_metric(run_id: str, field: str) -> str:
    path = EXPERIMENTS / run_id / "metrics.json"
    if not path.exists():
        return f"[pending:{run_id}.{field}]"
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if data.get("status") == "skipped":
        return f"[skipped: {data.get('reason', 'n/a')}]"
    value = data.get(field, "?")
    if field in ("cer", "wer") and isinstance(value, (int, float)):
        pct = f"{value * 100:.2f}%"
        if data.get("protocol") == "smoke_synthetic":
            return f"{pct} [SMOKE—not primary]"
        return pct
    return str(value)


def fill(text: str) -> str:
    def repl(m: re.Match) -> str:
        return load_metric(m.group(1), m.group(2))

    return re.sub(r"\{\{METRIC:([^.]+)\.([^}]+)\}\}", repl, text)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(fill(text))
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"


def main() -> None:
    doc = Document()
    doc.add_heading("Chapter 4 — Implementation and Results", 0)

    doc.add_heading("4.1 Purpose and contribution of the chapter", 1)
    add_para(
        doc,
        "Where Chapter 3 defined the research design, the present chapter reports how that "
        "design was realised and what was observed. The contribution is best understood as "
        "an aligned research-and-application study rather than as the invention of a new HTR "
        "architecture. In concrete terms, the chapter offers four interlocking outcomes: a "
        "reproducible English IAM training and evaluation workflow; a fully logged TrOCR "
        "fine-tuning experiment on the official training and validation partitions; a "
        "controlled comparison against a CRNN+CTC baseline on the identical full test set of "
        "2,915 lines; and an interactive demonstration that places the recogniser in an "
        "applied setting. Comparative reference to published IAM CER figures is included for "
        "contextual scale, without any claim of protocol-identical superiority.",
    )

    doc.add_heading("4.2 Implementation overview", 1)
    add_para(
        doc,
        "The implementation follows the methodological commitments of Chapter 3. Dataset "
        "handling loads line images and transcriptions for training and evaluation. Model "
        "utilities load either the public TrOCR checkpoint or a locally fine-tuned directory "
        "and generate transcriptions. A dedicated training routine adapts TrOCR with "
        "cross-entropy, AdamW, and early stopping on validation CER. A separate CRNN+CTC "
        "module supplies the neural baseline. Shared metric code computes case-sensitive CER "
        "and WER after whitespace normalisation. Evaluation runners execute the experiment "
        "matrix, while a Gradio interface exposes recognition for interactive demonstration. "
        "Detailed module listings are placed in the appendix so that the chapter body may "
        "remain focused on scientific substance rather than software documentation.",
    )
    doc.add_heading("Table 4.1 — Principal implementation components", 2)
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    for i, row in enumerate(
        [
            ("Component", "Role in the study"),
            ("Dataset I/O", "Provides line images and transcriptions for train/eval"),
            ("TrOCR loading and generation", "Instantiates Hub or local checkpoints for inference"),
            ("TrOCR fine-tuning", "Adapts the pretrained model on IAM train/val"),
            ("CRNN+CTC baseline", "Trains and evaluates the neural reference system"),
            ("Metrics", "Computes CER and WER under a shared definition"),
            ("Evaluation runners", "Executes the designed experiment matrix"),
            ("Demonstration interface", "Supports interactive English line transcription"),
        ]
    ):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val

    doc.add_heading("4.3 Experimental procedure", 1)
    doc.add_heading("4.3.1 TrOCR fine-tuning", 2)
    add_para(
        doc,
        "Fine-tuning began from microsoft/trocr-base-handwritten and used the full IAM "
        "training and validation sets (6,482 and 976 lines). Optimisation employed AdamW "
        "with a learning rate of 5×10⁻⁶, a micro-batch size of 1, and gradient accumulation "
        "of 8, yielding an effective batch size of 8. A maximum of eight epochs was planned, "
        "with early-stopping patience of two validations. Training ran on CUDA and required "
        "approximately 88.7 hours of wall-clock time.",
    )
    add_para(
        doc,
        "Under this schedule, validation CER reached its best value of 0.0137 after the "
        "first epoch and then deteriorated (0.0162 in epoch 2; 0.0187 in epoch 3). Early "
        "stopping therefore terminated training after three epochs, and the exported model "
        "corresponds to the best validation checkpoint rather than the final epoch weights. "
        "This behaviour is material to later interpretation: the adaptation run was completed "
        "and logged, yet it did not exhibit sustained validation improvement.",
    )

    doc.add_heading("4.3.2 CRNN+CTC training", 2)
    add_para(
        doc,
        "The CRNN baseline was trained from scratch on the same IAM partitions with learning "
        "rate 1×10⁻³, batch size 8, and fixed input geometry of 32×128. Eight epochs were "
        "executed under early-stopping patience, with a character set of size 79. Best "
        "validation CER remained high (approximately 0.80), already signalling limited "
        "competitiveness relative to pretrained TrOCR before test evaluation.",
    )

    doc.add_heading("4.3.3 Evaluation matrix", 2)
    add_para(
        doc,
        "Evaluation followed the design of Chapter 3. Primary claims rest on pretrained "
        "TrOCR (E1), fine-tuned TrOCR (E2), and CRNN+CTC (E3), each assessed on the full IAM "
        "test split. Tesseract (E4) remains optional and is reported as skipped when the "
        "binary is unavailable. The demonstration run (E5) supports the application section "
        "and is excluded from primary CER evidence.",
    )

    doc.add_heading("4.4 Results", 1)
    add_para(doc, "Table 4.2 reports the primary measured outcomes.")
    doc.add_heading("Table 4.2 — Measured results on the full IAM English test set", 2)
    mtable = doc.add_table(rows=4, cols=4)
    mtable.style = "Table Grid"
    for i, row in enumerate(
        [
            ("Method", "CER", "WER", "n"),
            (
                "TrOCR pretrained",
                "{{METRIC:iam_trocr_handwritten.cer}}",
                "{{METRIC:iam_trocr_handwritten.wer}}",
                "{{METRIC:iam_trocr_handwritten.n_samples}}",
            ),
            (
                "TrOCR fine-tuned",
                "{{METRIC:iam_trocr_finetuned.cer}}",
                "{{METRIC:iam_trocr_finetuned.wer}}",
                "{{METRIC:iam_trocr_finetuned.n_samples}}",
            ),
            (
                "CRNN+CTC",
                "{{METRIC:iam_crnn.cer}}",
                "{{METRIC:iam_crnn.wer}}",
                "{{METRIC:iam_crnn.n_samples}}",
            ),
        ]
    ):
        for j, val in enumerate(row):
            mtable.rows[i].cells[j].text = fill(val)

    add_para(
        doc,
        "The pretrained transformer attains a CER of {{METRIC:iam_trocr_handwritten.cer}} "
        "and a WER of {{METRIC:iam_trocr_handwritten.wer}}. After fine-tuning, CER and WER "
        "become {{METRIC:iam_trocr_finetuned.cer}} and {{METRIC:iam_trocr_finetuned.wer}}, "
        "respectively. The CRNN+CTC baseline yields a CER of {{METRIC:iam_crnn.cer}} and a "
        "WER of {{METRIC:iam_crnn.wer}} on the same population.",
    )
    add_para(
        doc,
        "Tesseract status for this study is: {{METRIC:iam_tesseract.cer}}. No classical OCR "
        "score is claimed when the run is skipped. Demonstration metrics, where present, are "
        "treated as qualitative application evidence only.",
    )

    doc.add_heading("Figure 4.1 — CER comparison", 2)
    if FIGURE.exists():
        doc.add_picture(str(FIGURE), width=Inches(5.8))
    add_para(
        doc,
        "Measured CER values for the systems evaluated in this thesis, with selected "
        "literature figures shown separately for contextual comparison.",
    )

    doc.add_heading("Table 4.3 — Literature CER figures cited for context", 2)
    ltable = doc.add_table(rows=4, cols=3)
    ltable.style = "Table Grid"
    for i, row in enumerate(
        [
            ("Method", "CER (IAM)", "Source"),
            ("AttentionHTR", "6.50%", "Kass & Vats, 2022"),
            ("Light Transformer", "5.70%", "Barrère et al., 2022"),
            ("GFCN", "7.99%", "Coquenet et al., 2020"),
        ]
    ):
        for j, val in enumerate(row):
            ltable.rows[i].cells[j].text = val
    add_para(
        doc,
        "These published values help the reader locate the present scores within the broader "
        "IAM literature. They are not re-implemented under the preprocessing and split "
        "conventions of this repository, and the thesis therefore does not assert a matched "
        "state-of-the-art ranking.",
    )

    doc.add_heading("4.5 Discussion and analysis", 1)
    doc.add_heading("4.5.1 Pretrained TrOCR as the principal empirical result", 2)
    add_para(
        doc,
        "The pretrained TrOCR evaluation constitutes the central positive finding of the "
        "study. A CER of {{METRIC:iam_trocr_handwritten.cer}} on 2,915 previously unseen IAM "
        "lines indicates that a publicly released transformer HTR checkpoint can deliver "
        "strong line-level recognition when assessed under a complete and transparent local "
        "protocol. For a master’s investigation concerned with transfer learning rather than "
        "architectural invention, this result is scientifically meaningful: it shows what "
        "current pretrained models already achieve before any student-led adaptation is "
        "attempted.",
    )

    doc.add_heading("4.5.2 Interpretation of the fine-tuning outcome", 2)
    add_para(
        doc,
        "Fine-tuning did not improve test recognition. The adapted model’s CER of "
        "{{METRIC:iam_trocr_finetuned.cer}} is slightly higher than that of the Hub "
        "checkpoint. Several considerations make this outcome intelligible. First, the "
        "starting checkpoint is already specialised for English handwriting, so residual "
        "headroom on IAM is limited. Second, validation CER worsened after the first epoch, "
        "and early stopping rightly preferred the earlier checkpoint; continued training "
        "under the logged schedule was not justified by the validation signal. Third, the "
        "optimisation used a small micro-batch with gradient accumulation, which may be "
        "adequate for a feasibility study yet suboptimal for extracting further gains. The "
        "appropriate academic reading is therefore that fine-tuning was executed and "
        "measured, and that its effect under these conditions was neutral to slightly "
        "adverse—not that adaptation is universally unhelpful, nor that the experiment "
        "failed to run.",
    )

    doc.add_heading("4.5.3 The CRNN+CTC baseline in perspective", 2)
    add_para(
        doc,
        "The CRNN+CTC system’s CER of {{METRIC:iam_crnn.cer}} confirms a substantial gap "
        "relative to TrOCR under matched data and metrics. A compact CNN–BiLSTM–CTC model "
        "trained from scratch with constrained geometry is not expected, in this setting, "
        "to rival a large pretrained encoder–decoder. The baseline nevertheless fulfils its "
        "methodological role: it anchors the transformer result against a familiar neural "
        "alternative and prevents the discussion from resting solely on a single model family.",
    )

    doc.add_heading("4.5.4 Error behaviour and qualitative observations", 2)
    add_para(
        doc,
        "Inspection of high-CER predictions reveals recurring difficulties with very short "
        "references, punctuation-dominated or dash-like strings, and rare or hyphenated "
        "tokens. Such lines can inflate per-example CER even when corpus-level error remains "
        "comparatively low. These observations caution against over-interpreting isolated "
        "failures and reinforce the primacy of aggregate CER/WER on the full test partition.",
    )

    doc.add_heading("4.5.5 Limitations affecting interpretation", 2)
    add_para(
        doc,
        "Interpretation remains bounded by the study’s design. The work addresses English "
        "line-level IAM only. Literature comparisons are contextual rather than "
        "protocol-matched. Tesseract may be unavailable in a given environment. The "
        "fine-tuning schedule was short and early-stopped, and should not be mistaken for "
        "an exhaustive hyperparameter search. Demonstration scores are illustrative and do "
        "not substitute for full-test evidence.",
    )

    doc.add_heading("4.6 Application prototype", 1)
    add_para(
        doc,
        "Beyond offline evaluation, the study includes a Gradio-based demonstration for "
        "interactive English line transcription, optionally with reference-based CER when "
        "ground truth is supplied. The prototype is an applied complement to the experimental "
        "programme: it shows that the recognition stack can be used by a non-specialist "
        "interface, while leaving scientific claims dependent on the full-test runs reported "
        "above.",
    )

    doc.add_heading("4.7 Reproducibility statement", 1)
    add_para(
        doc,
        "All quantitative statements in this chapter are intended to be recoverable from the "
        "logged experiment artefacts of the project. Training hyperparameters are taken from "
        "the fine-tuning and CRNN training logs; test metrics are taken from the corresponding "
        "evaluation records. Reproduction follows the project’s documented GPU protocol and "
        "evaluation commands. In this way, Chapter 4 remains answerable to the methodological "
        "design of Chapter 3 and prepares the reflective synthesis developed in Chapter 5.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
